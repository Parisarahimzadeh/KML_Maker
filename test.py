# -*- coding: utf-8 -*-
import os, re
import simplekml
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from xml.etree import ElementTree as ET

CSS_COLORS = {
    "black": (0,0,0), "white": (255,255,255), "red": (255,0,0),
    "green": (0,128,0), "blue": (0,0,255), "yellow": (255,255,0),
    "cyan": (0,255,255), "magenta": (255,0,255), "gray": (128,128,128),
    "darkgray": (169,169,169), "lightgray": (211,211,211),
    "orange": (255,165,0), "gold": (255,215,0), "purple": (128,0,128),
    "navy": (0,0,128), "teal": (0,128,128), "maroon": (128,0,0),
    "olive": (128,128,0), "lime": (0,255,0), "aqua": (0,255,255),
    "fuchsia": (255,0,255), "silver": (192,192,192),
    "brown": (165,42,42), "darkgreen": (0,100,0), "darkred": (139,0,0),
    "darkblue": (0,0,139), "deepskyblue": (0,191,255),
}

def parse_color_to_kml(value: str):
    if not value: return simplekml.Color.rgb(255,0,0)
    s = re.match(r"[a-z#0-9, ]+", value.strip().lower())
    s = s.group(0) if s else value.strip().lower()
    if s.startswith("#") and len(s)==7:
        r=int(s[1:3],16); g=int(s[3:5],16); b=int(s[5:7],16)
        return simplekml.Color.rgb(r,g,b,255)
    if "," in s:
        r,g,b = [int(x) for x in s.split(",")]
        return simplekml.Color.rgb(r,g,b,255)
    if s in CSS_COLORS:
        r,g,b = CSS_COLORS[s]
        return simplekml.Color.rgb(r,g,b,255)
    return simplekml.Color.rgb(255,0,0,255)

def _apply_tint(rgb_hex, tint):
    if tint in (None, 0): return rgb_hex
    r=int(rgb_hex[0:2],16); g=int(rgb_hex[2:4],16); b=int(rgb_hex[4:6],16)
    def adj(c):
        c/=255.0
        c = c*(1+tint) if tint < 0 else c + (1-c)*tint
        return max(0,min(1,c))
    r=int(adj(r)*255); g=int(adj(g)*255); b=int(adj(b)*255)
    return f"{r:02X}{g:02X}{b:02X}"

def _office_theme_map(wb):
    try:
        ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        root=ET.fromstring(wb.loaded_theme); sch=root.find('.//a:clrScheme',ns)
        names=['dk1','lt1','dk2','lt2','accent1','accent2','accent3','accent4','accent5','accent6']
        out=[]
        for n in names:
            node=sch.find(f'a:{n}',ns)
            srgb=node.find('a:srgbClr',ns)
            if srgb is not None: out.append(srgb.attrib.get('val','000000').upper())
            else:
                sysc=node.find('a:sysClr',ns)
                out.append((sysc.attrib.get('lastClr','000000').upper()) if sysc is not None else '000000')
        return out
    except: return ['000000','FFFFFF','44546A','E7E6E6','4472C4','ED7D31','A5A5A5','FFC000','5B9BD5','70AD47']

def _theme_idx_to_hex(wb, idx, tint):
    scheme=_office_theme_map(wb)
    if isinstance(idx,int) and 0<=idx<len(scheme):
        return _apply_tint(scheme[idx], tint)
    return None

def excel_bg_hex(cell, wb):
    """فقط وقتی fill='solid' باشد رنگ پس‌زمینه را برگردان؛ در غیر این صورت None."""
    fill = getattr(cell, "fill", None)
    if not fill or fill.patternType != "solid":
        return None
    color = fill.fgColor or fill.start_color
    if not color: return None

    rgb = getattr(color, "rgb", None)
    if isinstance(rgb, str) and len(rgb) in (6,8):
        core = rgb[2:] if len(rgb)==8 else rgb
        return f"#{core.upper()}"

    if getattr(color,"type",None)=="indexed":
        idx=getattr(color,"indexed",None)
        if isinstance(idx,int) and 0<=idx<len(COLOR_INDEX):
            argb=COLOR_INDEX[idx]
            if isinstance(argb,str) and len(argb) in (6,8):
                core = argb[2:] if len(argb)==8 else argb
                return f"#{core.upper()}"

    if getattr(color,"type",None)=="theme":
        idx=getattr(color,"theme",None); tint=getattr(color,"tint",None)
        base=_theme_idx_to_hex(wb, idx, tint)
        if base: return f"#{base}"
    return None

def text_color_for_bg(bg_hex):
    if not bg_hex: return ""
    try:
        h=bg_hex.lstrip("#"); r=int(h[0:2],16); g=int(h[2:4],16); b=int(h[4:6],16)
        l=0.299*r+0.587*g+0.114*b
        return "color:white;" if l<128 else ""
    except: return ""

root = tk.Tk(); root.withdraw()
excel_file = filedialog.askopenfilename(title="choosing file", filetypes=[("Excel files","*.xlsx *.xls")])
if not excel_file:
    messagebox.showerror("Error","Please select one file"); raise SystemExit()

base_path = os.path.dirname(excel_file)
out_dir = os.path.join(base_path, "KML_Output"); os.makedirs(out_dir, exist_ok=True)

wb = load_workbook(excel_file, data_only=True)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=False))
    if not rows: print(f"⛔ sheet '{sheet_name}' rejected. Empty sheet"); continue

    headers = [(str(c.value).strip() if c.value is not None else "") for c in rows[0]]

    def find_col(key):
        k=key.lower()
        for i,h in enumerate(headers):
            if k in h.lower(): return i
        return None

    lat_idx=find_col("lat"); lon_idx=find_col("lon"); dot_idx=find_col("dotcolor")
    if lat_idx is None or lon_idx is None:
        print(f"⛔ sheet '{sheet_name}' rejected. No Lat/Lon"); continue

    kml = simplekml.Kml()
    circle_icon = "http://maps.google.com/mapfiles/kml/shapes/placemark_circle.png"

    for r, row_cells in enumerate(rows[1:], start=2):
        vals=[c.value if c is not None else None for c in row_cells]
        lat = vals[lat_idx] if lat_idx<len(vals) else None
        lon = vals[lon_idx] if lon_idx<len(vals) else None
        if lat is None or lon is None: continue

        dot_val = str(vals[dot_idx]) if (dot_idx is not None and dot_idx<len(vals) and vals[dot_idx] is not None) else None

        rows_html=[]
        for c_idx, hdr in enumerate(headers):
            cell = row_cells[c_idx] if c_idx<len(row_cells) else None
            bg = excel_bg_hex(cell, wb) if cell is not None else None

            # فقط BAPath/BHPath را اگر سفید/بی‌رنگ بودند مشکی کن
            hdr_norm = re.sub(r'[\s_\-]+','', (hdr or '')).lower()
            if (hdr_norm.startswith("bapath") or hdr_norm.startswith("bhpath")) and (bg is None or bg.upper()=="#FFFFFF"):
                bg = "#000000"

            txt = text_color_for_bg(bg)
            style = (f"background-color:{bg};" if bg else "") + txt
            val = "" if vals[c_idx] is None else str(vals[c_idx])
            border = "border:1px solid #bbb;" if bg else "border:1px solid #ddd;"

            rows_html.append(
                "<tr>"
                f"<th style='border:1px solid #ccc;padding:6px;background:#f2f2f2;text-align:left;white-space:nowrap;min-width:140px'>{hdr}</th>"
                f"<td style='{border}padding:6px;{style}min-width:220px'>{val}</td>"
                "</tr>"
            )

        # امضا
        rows_html.append(
            "<tr>"
            "<th style='border:1px solid #ccc;padding:6px;background:#f2f2f2;text-align:left;white-space:nowrap'>Creator</th>"
            "<td style='border:1px solid #ddd;padding:6px;min-width:220px'>"
            "<a href='https://www.linkedin.com/in/parisa-rahimzadeh' target='_blank' style='color:#2E86C1;text-decoration:none;'>Created by Parisa Rahimzadeh</a>"
            "</td>"
            "</tr>"
        )

        table_html = "<table style='border-collapse:collapse;font-family:Arial;font-size:12px;width:520px;table-layout:auto'>" + "".join(rows_html) + "</table>"
        point_name = f"{sheet_name}_Row{r}"
        desc = "<![CDATA[" + f"<h3 style='color:#2E86C1;margin:0;padding:0;'>{point_name}</h3><hr>" + table_html + "]]>"

        pnt = kml.newpoint(name=point_name, coords=[(float(lon), float(lat))])
        pnt.description = desc
        pnt.style.labelstyle.scale = 0.9
        pnt.style.iconstyle.icon.href = circle_icon
        pnt.style.iconstyle.scale = 1.0
        pnt.style.iconstyle.color = parse_color_to_kml(dot_val) if dot_val else simplekml.Color.rgb(255,0,0,255)

    out_file = os.path.join(out_dir, f"{sheet_name}.kml")
    kml.save(out_file)
    print(f"✅ created: {out_file}")

messagebox.showinfo("Done", f"All KML files created in \n{out_dir}\n")
